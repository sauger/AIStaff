from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from app.cabinet.errors import CabinetError

TEXT_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".log",
}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx"}


def can_apply_replacements(filename: str) -> bool:
    suffix = Path(filename).suffix.lower()
    return suffix in TEXT_SUFFIXES or suffix in DOCX_SUFFIXES or suffix in XLSX_SUFFIXES


def apply_replacements(filename: str, data: bytes, replacements: dict[str, str]) -> bytes:
    if not replacements:
        return data
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in TEXT_SUFFIXES:
            return _replace_text(data, replacements)
        if suffix in DOCX_SUFFIXES:
            return _replace_docx(data, replacements)
        if suffix in XLSX_SUFFIXES:
            return _replace_xlsx(data, replacements)
    except CabinetError:
        raise
    except Exception as exc:
        raise CabinetError(
            "CABINET_TEMPLATE_UNPROCESSABLE",
            f"无法处理模板「{filename}」：{exc}。原模板未被覆盖，也没有生成新文件。",
            details={"filename": filename, "reason": type(exc).__name__},
        ) from exc
    raise CabinetError(
        "CABINET_TEMPLATE_UNPROCESSABLE",
        (
            f"无法按模板生产「{filename}」：运行时不能改写该格式。"
            "请改用可编辑的办公文档，或不要要求改内容。原模板仍在，没有生成新文件。"
        ),
        details={"filename": filename, "suffix": suffix or "(none)"},
    )


def _replace_text(data: bytes, replacements: dict[str, str]) -> bytes:
    text = None
    used = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = data.decode(encoding)
            used = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None or used is None:
        raise CabinetError(
            "CABINET_TEMPLATE_UNPROCESSABLE",
            "无法按模板生产：文本文件解码失败。原模板仍在，没有生成新文件。",
        )
    for old, new in replacements.items():
        text = text.replace(old, new)
    if used == "utf-8-sig":
        return text.encode("utf-8-sig")
    return text.encode(used)


def _replace_docx(data: bytes, replacements: dict[str, str]) -> bytes:
    from docx import Document

    document = Document(BytesIO(data))
    for paragraph in document.paragraphs:
        _replace_paragraph(paragraph, replacements)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, replacements)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _replace_paragraph(paragraph: object, replacements: dict[str, str]) -> None:
    text = getattr(paragraph, "text", "") or ""
    updated = text
    for old, new in replacements.items():
        updated = updated.replace(old, new)
    if updated == text:
        return
    runs = getattr(paragraph, "runs", [])
    if runs:
        runs[0].text = updated
        for run in runs[1:]:
            run.text = ""
        return
    add_run = getattr(paragraph, "add_run", None)
    if callable(add_run):
        add_run(updated)


def _replace_xlsx(data: bytes, replacements: dict[str, str]) -> bytes:
    try:
        source = zipfile.ZipFile(BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise CabinetError(
            "CABINET_TEMPLATE_UNPROCESSABLE",
            "无法按模板生产：Excel 文件损坏或不是有效的 .xlsx。原模板仍在，没有生成新文件。",
        ) from exc
    escaped = {old: escape(new) for old, new in replacements.items()}
    output = BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as dest:
        touched = False
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.endswith("sharedStrings.xml") or "/worksheets/" in info.filename:
                try:
                    text = payload.decode("utf-8")
                except UnicodeDecodeError as exc:
                    raise CabinetError(
                        "CABINET_TEMPLATE_UNPROCESSABLE",
                        "无法按模板生产：无法读取 Excel 内部文本。原模板仍在，没有生成新文件。",
                    ) from exc
                updated = text
                for old, new in escaped.items():
                    updated = updated.replace(old, new)
                if updated != text:
                    touched = True
                payload = updated.encode("utf-8")
            dest.writestr(info, payload)
        if not touched and replacements:
            raise CabinetError(
                "CABINET_TEMPLATE_UNPROCESSABLE",
                "无法按模板生产：该 Excel 中没有可替换的共享文本。原模板仍在，没有生成新文件。",
            )
    return output.getvalue()
